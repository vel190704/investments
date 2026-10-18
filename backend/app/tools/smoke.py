"""Post-deploy smoke test against a DEPLOYED backend (run manually after each deploy; docs/DEPLOY.md §16).

    make smoke URL=https://<preview backend>

Two dummy accounts in the PREVIEW Supabase project (created with passwords, invited in the app) sign in. B uploads
a workbook with a freshly generated canary holding (unique symbol, name and note). A then calls every GET route
with B's ids, the Excel export, the status endpoints, the admin endpoints (403 expected) and some deliberate
errors. The canary must never appear and nothing may answer 5xx. Finally B's canary is removed again.

Environment (never in git; preview values only):
  SMOKE_TARGET=preview                    required; the script refuses otherwise
  SMOKE_SUPABASE_URL, SMOKE_SUPABASE_ANON_KEY   the PREVIEW project's URL and publishable key
  SMOKE_A_EMAIL, SMOKE_A_PASSWORD, SMOKE_B_EMAIL, SMOKE_B_PASSWORD
  SMOKE_PRODUCTION_SUPABASE_URL           optional; the script refuses if SMOKE_SUPABASE_URL equals it
`--dev` uses X-Dev-User headers instead of sign-in (only a local dev server accepts them; used by the tests).
"""

from __future__ import annotations

import argparse
import io
import os
import re
import secrets
import sys

import httpx

ERROR_PROBES = (("POST", "/api/journal/{iid}", {"why": "x", "change_view": "x", "horizon": "x"}),
                ("PUT", "/api/peers/{iid}", {"symbols": ["ABB"], "note": "x"}),
                ("DELETE", "/api/transactions/{txn_id}", None),
                ("GET", "/api/valuation/{iid}", None),
                ("GET", "/api/analysis/{iid}/preview", None))
PARAM = re.compile(r"{([^}:]+)(?::[^}]+)?}")


class SmokeError(SystemExit):
    pass


def supabase_token(sb_url: str, anon_key: str, email: str, password: str, http: httpx.Client) -> str:
    r = http.post(f"{sb_url.rstrip('/')}/auth/v1/token?grant_type=password",
                  headers={"apikey": anon_key, "Content-Type": "application/json"},
                  json={"email": email, "password": password}, timeout=30)
    if r.status_code != 200:
        raise SmokeError(f"sign-in failed for an account ({r.status_code}); check the preview users and passwords")
    return r.json()["access_token"]


def canary_workbook(token: str) -> tuple[bytes, dict]:
    import openpyxl

    from app.importer.validate import TXN_COLUMNS
    from app.services.template import FALLBACK_COLUMNS, PLAN_COLUMNS

    c = {"symbol": "ZQ" + token[:6].upper(), "name": f"Smoke Canary {token}", "note": f"smoke-note-{token}"}
    wb = openpyxl.Workbook()
    wb.active.title = "README"
    t = wb.create_sheet("transactions")
    t.append(TXN_COLUMNS)
    t.append(["2026-01-02", "OPENING", "STOCK", c["name"], c["symbol"], None, None, 3, 111, 333, 0, "Smoke",
              c["note"]])
    wb.create_sheet("holdings_fallback").append(FALLBACK_COLUMNS)
    wb.create_sheet("monthly_plan").append(PLAN_COLUMNS)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), c


def empty_workbook() -> bytes:
    from app.services.template import build_template

    return build_template()


def _flat_get_routes(app) -> list[str]:
    """Every GET route of this codebase (the deployed version should match the checked-out commit)."""
    from fastapi.routing import APIRoute

    out = []

    def walk(routes, prefix=""):
        for r in routes:
            if isinstance(r, APIRoute):
                if "GET" in r.methods:
                    out.append(prefix + r.path)
            elif hasattr(r, "original_router"):
                walk(r.original_router.routes, prefix + r.include_context.prefix)
    walk(app.routes)
    return sorted(set(out))


def _text(r: httpx.Response) -> str:
    if "spreadsheetml" in r.headers.get("content-type", ""):
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True)
        return "\n".join(str(v) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row if v)
    return r.text


def run(base: str, http: httpx.Client, ha: dict, hb: dict, out=print) -> dict:
    """The smoke test proper. ha/hb: request headers that authenticate A and B. Returns the report."""
    from app.main import app

    base = base.rstrip("/")
    for h, who in ((ha, "A"), (hb, "B")):
        r = http.get(f"{base}/api/me", headers=h)
        if r.status_code != 200:
            raise SmokeError(f"account {who} cannot use the app ({r.status_code}); is it invited?")
    token = secrets.token_hex(6)
    book, c = canary_workbook(token)
    canaries = (c["symbol"], c["name"], c["note"])
    r = http.post(f"{base}/api/import/preview", headers=hb, files={"file": ("smoke.xlsx", book)})
    r.raise_for_status()
    r = http.post(f"{base}/api/import/confirm", headers=hb, json={"token": r.json()["token"]})
    r.raise_for_status()
    txns = http.get(f"{base}/api/transactions", headers=hb).json()
    mine = [t for t in txns if t.get("name") == c["name"]]
    if not mine:
        raise SmokeError("B's canary holding was not imported")
    ids = {"iid": mine[0]["instrument_id"], "txn_id": mine[0]["id"], "cid": 999_999, "did": 999_999,
           "aid": 999_999, "user_id": 999_999}
    if not any(x in http.get(f"{base}/api/transactions", headers=hb).text for x in canaries):
        raise SmokeError("B cannot see their own canary: the test would prove nothing")

    problems, checked = [], 0

    def check(method, path, resp):
        nonlocal checked
        checked += 1
        body = _text(resp)
        if resp.status_code >= 500:
            problems.append(f"{method} {path}: {resp.status_code}")
        for x in canaries:
            if x in body:
                problems.append(f"{method} {path}: B's canary visible to A")

    def url(path):
        return base + PARAM.sub(lambda m: str(ids.get(m.group(1), {"section": "drift", "metric": "roce",
                                                                     "field": "sector", "key": "repo_rate",
                                                                     "sector": "Construction",
                                                                     "email": "nobody@example.com"}.get(m.group(1), 1))),
                                path)
    try:
        for path in _flat_get_routes(app):
            if path == "/api/health":
                continue
            params = {"a": "2026-01-01", "b": "live"} if path == "/api/compare" else None
            check("GET", path, http.get(url(path), headers=ha, params=params))
        for method, path, body in ERROR_PROBES:
            check(method, path, http.request(method, url(path), headers=ha, json=body))
        r = http.get(f"{base}/api/admin/users", headers=ha)
        if r.status_code not in (403, 200):
            problems.append(f"GET /api/admin/users: unexpected {r.status_code}")
        if r.status_code == 200:
            problems.append("account A has admin rights: use a normal account for the smoke test")
    finally:
        r = http.post(f"{base}/api/import/preview", headers=hb, files={"file": ("empty.xlsx", empty_workbook())})
        if r.status_code == 200:
            http.post(f"{base}/api/import/confirm", headers=hb, json={"token": r.json()["token"]})
        left = any(x in http.get(f"{base}/api/transactions", headers=hb).text for x in canaries)
        if left:
            problems.append("cleanup: B's canary is still there (remove it on B's Transactions page)")
    report = {"checked": checked, "problems": problems}
    out(f"smoke: {checked} requests as A checked for B's canary; "
        + ("PASS" if not problems else f"FAIL ({len(problems)})"))
    for p in problems:
        out("  - " + p)
    return report


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="Cross-user smoke test against a deployed (preview) backend.")
    ap.add_argument("--url", required=True, help="backend base URL, e.g. https://investments-api-preview.onrender.com")
    ap.add_argument("--dev", action="store_true", help="local dev server only: X-Dev-User headers, no sign-in")
    args = ap.parse_args(argv)
    with httpx.Client(timeout=120, follow_redirects=False) as http:
        if args.dev:
            ha, hb = {"X-Dev-User": "smoke-a@example.com"}, {"X-Dev-User": "smoke-b@example.com"}
        else:
            env = os.environ
            if env.get("SMOKE_TARGET") != "preview":
                raise SmokeError("refusing: set SMOKE_TARGET=preview (the smoke test writes test data; never production)")
            sb = env.get("SMOKE_SUPABASE_URL", "")
            if sb and sb.rstrip("/") == env.get("SMOKE_PRODUCTION_SUPABASE_URL", "").rstrip("/"):
                raise SmokeError("refusing: SMOKE_SUPABASE_URL is the production project")
            missing = [k for k in ("SMOKE_SUPABASE_URL", "SMOKE_SUPABASE_ANON_KEY", "SMOKE_A_EMAIL", "SMOKE_A_PASSWORD",
                                   "SMOKE_B_EMAIL", "SMOKE_B_PASSWORD") if not env.get(k)]
            if missing:
                raise SmokeError(f"missing environment variables: {', '.join(missing)}")
            ha = {"Authorization": "Bearer " + supabase_token(sb, env["SMOKE_SUPABASE_ANON_KEY"], env["SMOKE_A_EMAIL"],
                                                              env["SMOKE_A_PASSWORD"], http)}
            hb = {"Authorization": "Bearer " + supabase_token(sb, env["SMOKE_SUPABASE_ANON_KEY"], env["SMOKE_B_EMAIL"],
                                                              env["SMOKE_B_PASSWORD"], http)}
        report = run(args.url, http, ha, hb)
    return 0 if not report["problems"] else 1


if __name__ == "__main__":
    sys.exit(main())
