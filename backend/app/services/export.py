"""Export the DB back to a workbook in exactly the input layout (valid backup + valid import)."""

from __future__ import annotations

import io
from decimal import Decimal

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import repository as repo
from app.importer.validate import TXN_COLUMNS
from app.models import FallbackHolding, Instrument
from app.money import paise_to_rupees
from app.timeutil import today_ist

FALLBACK_COLUMNS = ["as_of_date", "platform", "asset_type", "name", "symbol", "quantity", "invested_amount",
                    "monthly_plan", "sector", "notes"]
PLAN_COLUMNS = ["platform", "asset_type", "name", "monthly_amount", "allocation_percent"]


def _num(d: Decimal | None):
    """Excel cell value for an exact decimal: int when whole, else Decimal (written as a number)."""
    if d is None:
        return None
    return int(d) if d == d.to_integral_value() else d


def _rupees(paise: int | None):
    return None if paise is None else _num(paise_to_rupees(paise))


def build_workbook(s: Session) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    ws.append(["Personal Investment Tracker (exported)", None])
    ws.append(["Exported on", today_ist().isoformat()])
    ws.append(["Note", "Contains workbook rows and in-app entries. Re-importing this file does not duplicate "
                       "in-app entries."])
    ws.append(["allocation_percent", "Stored as a FRACTION: 0.6 = 60%."])

    insts = {i.id: i for i in s.scalars(select(Instrument))}
    t = wb.create_sheet("transactions")
    t.append(TXN_COLUMNS)
    for x in repo.all_transactions(s):
        i = insts[x.instrument_id]
        t.append([
            x.date.isoformat(), x.type, i.asset_type, i.name, i.symbol if i.asset_type != "MF" else None,
            i.isin, i.scheme_code, _num(x.quantity), _num(x.price), _rupees(x.amount_paise),
            _rupees(x.charges_paise), x.platform, x.notes,
        ])

    f = wb.create_sheet("holdings_fallback")
    f.append(FALLBACK_COLUMNS)
    for fb in s.scalars(select(FallbackHolding).order_by(FallbackHolding.source_row)):
        i = insts[fb.instrument_id]
        f.append([fb.as_of_date.isoformat(), fb.platform, i.asset_type, i.name, i.symbol, _num(fb.quantity),
                  _rupees(fb.invested_paise), _rupees(fb.monthly_plan_paise), fb.sector, fb.notes])

    p = wb.create_sheet("monthly_plan")
    p.append(PLAN_COLUMNS)
    items = repo.plan_items(s)
    for it in items:
        p.append([it.platform, it.asset_type, it.name, _rupees(it.monthly_amount_paise), _num(it.allocation_fraction)])
    if items:
        p.append(["TOTAL", "ALL", "Total monthly investment", _rupees(sum(i.monthly_amount_paise for i in items)),
                  _num(sum((i.allocation_fraction for i in items), Decimal(0)))])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
