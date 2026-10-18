"""Blank workbook template for new users: the three input sheets with their headers, plus a README sheet.
No example rows in the input sheets (they would be imported as data); examples are in the README text only."""

from __future__ import annotations

import io

import openpyxl
from openpyxl.styles import Font

from app.importer.validate import TXN_COLUMNS

FALLBACK_COLUMNS = ["as_of_date", "platform", "asset_type", "name", "symbol", "quantity", "invested_amount",
                    "monthly_plan", "sector", "notes"]
PLAN_COLUMNS = ["platform", "asset_type", "name", "monthly_amount", "allocation_percent"]

README = [
    "Personal investment tracker: template",
    "",
    "1. transactions (main input): one row per buy, SIP, sell, dividend or opening balance.",
    "   date: YYYY-MM-DD · type: BUY, SIP, SELL, DIVIDEND or OPENING · asset_type: STOCK, ETF or MF",
    "   name: as you want to see it · symbol: NSE symbol for stocks/ETFs (e.g. LT) · scheme_code: AMFI code if known",
    "   quantity, price, amount (₹; for SIPs the amount is authoritative), charges (₹, optional), platform, notes",
    "   Example (write it in the transactions sheet, not here): 2026-10-03 | BUY | STOCK | Larsen & Toubro | LT | | |"
    " 1 | 3768 | 3768 | 0 | Groww |",
    "2. holdings_fallback (optional): current holdings when you don't have the full history.",
    "3. monthly_plan (optional): your monthly contribution split; allocation_percent is a fraction (0.6 = 60%).",
    "",
    "Upload it on the Import page: you see a preview first, and nothing changes until you confirm.",
    "Then confirm the proposed identifiers on the Mappings page (funds get NAVs only after you confirm a scheme code).",
    "This is a personal analytics tool, not investment advice.",
]


def build_template() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    for line in README:
        ws.append([line])
    ws["A1"].font = Font(bold=True)
    for title, cols in (("transactions", TXN_COLUMNS), ("holdings_fallback", FALLBACK_COLUMNS),
                        ("monthly_plan", PLAN_COLUMNS)):
        sh = wb.create_sheet(title)
        sh.append(cols)
        for c in sh[1]:
            c.font = Font(bold=True)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
