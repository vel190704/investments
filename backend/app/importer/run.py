"""Workbook -> validated rows -> idempotent sync into the DB, with a full report."""

from __future__ import annotations

import hashlib
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
from sqlalchemy.orm import Session

from app import repository as repo
from app.importer.validate import (
    FALLBACK_REQUIRED, PLAN_REQUIRED, TXN_REQUIRED, ParsedTxn, Problem, header_problems,
    normalise_name, parse_fallback, parse_plan, parse_txn, row_dict,
)
from app.models import FallbackHolding, IdentifierCandidate, ImportRun, Instrument, MonthlyPlanItem, Transaction
from app.services import changelog
from app.services.instruments import get_or_create
from app.money import format_inr
from app.timeutil import utcnow
from app.types import decimal_to_str


def _p(problems: list[Problem]) -> list[dict]:
    return [{"row": p.row, "field": p.field, "message": p.message} for p in problems]


def _rows(ws) -> tuple[list, list[tuple[int, tuple]]]:
    it = ws.iter_rows(values_only=True)
    header = list(next(it, []) or [])
    rows = []
    for i, values in enumerate(it, start=2):
        if values is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
            continue  # fully blank line: not a data row
        rows.append((i, values))
    return header, rows


def identity(d, ttype: str, key: str, quantity, price, amount_paise: int, charges_paise: int,
             platform: str | None) -> str:
    """Business identity of a transaction (notes and source excluded)."""
    parts = [d.isoformat(), ttype, key,
             decimal_to_str(quantity) if quantity is not None else "",
             decimal_to_str(price) if price is not None else "",
             str(amount_paise), str(charges_paise), platform or ""]
    return hashlib.sha256("|".join(parts).encode()).hexdigest()[:40]


def row_identity(t: ParsedTxn) -> str:
    return identity(t.date, t.type, t.instrument_key, t.quantity, t.price, t.amount_paise, t.charges_paise,
                    t.platform)


def db_identity(t: Transaction) -> str:
    return identity(t.date, t.type, t.instrument.instrument_key, t.quantity, t.price, t.amount_paise,
                    t.charges_paise, t.platform)


def check_oversell(txns: list[ParsedTxn], context: list[ParsedTxn] = ()) -> tuple[list[ParsedTxn], list[Problem]]:
    """Drop SELL rows that exceed the quantity held at that point (chronological replay).

    `context` rows (in-app entries already in the DB) take part in the replay but are never dropped.
    """
    from app.calc.positions import TYPE_ORDER

    by_inst: dict[str, list[ParsedTxn]] = defaultdict(list)
    for t in [*txns, *context]:
        by_inst[t.instrument_key].append(t)
    bad: set[int] = set()
    problems = []
    for rows in by_inst.values():
        held = Decimal(0)
        for t in sorted(rows, key=lambda t: (t.date, TYPE_ORDER[t.type], t.row)):
            if t.type in ("OPENING", "BUY", "SIP"):
                held += t.quantity
            elif t.type == "SELL":
                if t.quantity > held and t.row > 0:
                    bad.add(t.row)
                    problems.append(Problem(t.row, f"SELL of {t.quantity} exceeds quantity held ({held}) on {t.date}",
                                            "quantity"))
                else:
                    held -= t.quantity
    return [t for t in txns if t.row not in bad], problems


def _ui_context(session: Session) -> list[ParsedTxn]:
    from sqlalchemy import select

    out = []
    for t in session.scalars(select(Transaction).where(Transaction.source == "ui")):
        i = t.instrument
        out.append(ParsedTxn(row=0, date=t.date, type=t.type, asset_type=i.asset_type, name=i.name, symbol=i.symbol,
                             isin=None, scheme_code=None, quantity=t.quantity, price=t.price,
                             amount_paise=t.amount_paise, charges_paise=t.charges_paise, platform=t.platform,
                             notes=None))
    return out


def import_workbook(session: Session, path: str | Path, file_name: str | None = None) -> dict:
    path = Path(path)
    file_name = file_name or path.name
    data = path.read_bytes()
    report: dict = {
        "file": file_name, "sha256": hashlib.sha256(data).hexdigest(), "ok": False, "fatal": [],
        "transactions": {"rows_in_sheet": 0, "loaded": 0, "inserted": 0, "unchanged": 0, "updated_notes": 0,
                         "removed": 0, "matched_ui": [], "added_rows": [], "removed_rows": [], "pruning": "", "skipped": [], "errors": [], "warnings": [],
                         "by_asset_type": {}},
        "holdings_fallback": {"rows": 0, "used": [], "ignored": 0, "skipped": [], "errors": []},
        "monthly_plan": {"rows": 0, "total_paise": 0, "total_row_paise": None, "fraction_sum": "0",
                         "warnings": [], "errors": []},
        "instruments": {"created": [], "existing": 0},
    }
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # corrupt / not an xlsx
        report["fatal"].append(f"Could not open workbook: {e}")
        return _finish(session, report)

    for sheet, req in (("transactions", TXN_REQUIRED), ("holdings_fallback", FALLBACK_REQUIRED),
                       ("monthly_plan", PLAN_REQUIRED)):
        if sheet not in wb.sheetnames:
            if sheet == "transactions" or sheet == "monthly_plan":
                report["fatal"].append(f"Sheet '{sheet}' not found")
            continue
        report["fatal"] += header_problems(_rows(wb[sheet])[0], req, sheet)
    if report["fatal"]:
        return _finish(session, report)

    # ---- transactions
    tr = report["transactions"]
    header, rows = _rows(wb["transactions"])
    tr["rows_in_sheet"] = len(rows)
    parsed: list[ParsedTxn] = []
    errors: list[Problem] = []
    for row_no, values in rows:
        try:
            t, skip, errs = parse_txn(row_no, row_dict(header, values))
        except Exception as e:  # never crash on a bad row
            t, skip, errs = None, None, [Problem(row_no, f"unexpected error: {e}")]
        if skip:
            tr["skipped"].append({"row": row_no, "reason": skip})
        errors += errs
        if t:
            parsed.append(t)
            tr["warnings"] += [{"row": row_no, "message": w} for w in t.warnings]
    parsed, oversell = check_oversell(parsed, _ui_context(session))
    errors += oversell

    # ---- holdings_fallback
    fb = report["holdings_fallback"]
    fallbacks = []
    if "holdings_fallback" in wb.sheetnames:
        header_f, rows_f = _rows(wb["holdings_fallback"])
        fb["rows"] = len(rows_f)
        for row_no, values in rows_f:
            try:
                f, skip, errs = parse_fallback(row_no, row_dict(header_f, values))
            except Exception as e:
                f, skip, errs = None, None, [Problem(row_no, f"unexpected error: {e}")]
            if skip:
                fb["skipped"].append({"row": row_no, "reason": skip})
            fb["errors"] += _p(errs)
            if f:
                fallbacks.append(f)

    # ---- monthly_plan
    mp = report["monthly_plan"]
    header_p, rows_p = _rows(wb["monthly_plan"])
    plans = []
    for row_no, values in rows_p:
        r = row_dict(header_p, values)
        p, skip, errs = parse_plan(row_no, r)
        if skip == "TOTAL":
            try:
                from app.money import rupees_to_paise
                mp["total_row_paise"] = rupees_to_paise(r.get("monthly_amount"))
            except Exception:
                mp["warnings"].append(f"row {row_no}: TOTAL row amount is not a number")
        mp["errors"] += _p(errs)
        if p:
            plans.append(p)
    mp["rows"] = len(plans)
    mp["total_paise"] = sum(p.monthly_amount_paise for p in plans)
    frac_sum = sum((p.allocation_fraction for p in plans), Decimal(0))
    mp["fraction_sum"] = decimal_to_str(frac_sum)
    if plans and abs(frac_sum - 1) > Decimal("0.0001"):
        mp["warnings"].append(f"allocation_percent values sum to {frac_sum}, not 1 (100%)")
    if mp["total_row_paise"] is not None and mp["total_row_paise"] != mp["total_paise"]:
        mp["warnings"].append(f"items sum to {format_inr(mp['total_paise'])} but TOTAL row says "
                              f"{format_inr(mp['total_row_paise'])}")
    for p in plans:
        if mp["total_paise"] and abs(Decimal(p.monthly_amount_paise) / mp["total_paise"] - p.allocation_fraction) > Decimal("0.005"):
            mp["warnings"].append(f"row {p.row}: {p.name} amount share differs from allocation_percent")

    tr["errors"] = _p(errors)
    _sync(session, report, parsed, fallbacks, plans, file_name, has_errors=bool(errors))
    return _finish(session, report)


def _instrument(session: Session, cache: dict[str, Instrument], report: dict, key: str, asset_type: str,
                name: str, symbol: str | None, platform: str | None) -> Instrument:
    inst, created = get_or_create(session, cache, key, asset_type, name, symbol, platform)
    if created:
        report["instruments"]["created"].append(name)
    return inst


def _confirm_from_excel(session: Session, inst: Instrument, kind: str, value: str) -> None:
    """Identifiers typed into the Excel file are user-supplied: stored as CONFIRMED."""
    c = repo.find_candidate(session, inst.id, kind, value)
    if c is None:
        c = IdentifierCandidate(instrument_id=inst.id, kind=kind, value=value, label=inst.name,
                                source="Excel file", score=1.0, details={})
        session.add(c)
    c.status, c.decided_at = "CONFIRMED", utcnow()
    if kind == "ISIN":
        inst.isin = value
    else:
        inst.scheme_code = value


def _sync(session, report, parsed, fallbacks, plans, file_name, has_errors: bool) -> None:
    tr = report["transactions"]
    cache = repo.instruments_by_key(session)
    existing_count = len(cache)
    existing = repo.excel_row_keys(session)
    positions_before = changelog.position_states(session)
    # In-app (ui) rows by identity: an identical Excel row is recognised, not duplicated.
    ui_pool: Counter = Counter(db_identity(t) for t in repo.ui_transactions(session))
    added: list[Transaction] = []
    seen: set[str] = set()
    occurrences: Counter = Counter()
    by_type: Counter = Counter()
    for t in parsed:
        inst = _instrument(session, cache, report, t.instrument_key, t.asset_type, t.name, t.symbol, t.platform)
        if t.isin:
            _confirm_from_excel(session, inst, "ISIN", t.isin)
        if t.scheme_code:
            _confirm_from_excel(session, inst, "SCHEME_CODE", t.scheme_code)
        ident = row_identity(t)
        occurrences[ident] += 1
        key = f"{ident}#{occurrences[ident]}"
        seen.add(key)
        by_type[t.asset_type] += 1
        if key in existing:
            row = existing[key]
            tr["unchanged"] += 1
            if row.notes != t.notes:
                row.notes = t.notes
                tr["updated_notes"] += 1
            row.source_row, row.source_file = t.row, file_name
        elif ui_pool[ident] > 0:
            ui_pool[ident] -= 1
            tr["matched_ui"].append({"row": t.row, "message": "already present as an in-app entry; kept as is"})
        else:
            row = Transaction(
                instrument_id=inst.id, date=t.date, type=t.type, quantity=t.quantity, price=t.price,
                amount_paise=t.amount_paise, charges_paise=t.charges_paise, platform=t.platform, notes=t.notes,
                source="excel", row_key=key, source_file=file_name, source_row=t.row,
            )
            session.add(row)
            added.append(row)
            tr["inserted"] += 1
            tr["added_rows"].append({"row": t.row, "date": t.date.isoformat(), "type": t.type, "name": t.name,
                                     "quantity": decimal_to_str(t.quantity) if t.quantity is not None else None,
                                     "amount_paise": t.amount_paise})
    tr["loaded"] = len(parsed)
    tr["by_asset_type"] = dict(sorted(by_type.items()))

    stale = [row for k, row in existing.items() if k not in seen]
    tr["removed_rows"] = [{"id": r.id, "date": r.date.isoformat(), "type": r.type, "name": r.instrument.name,
                           "quantity": decimal_to_str(r.quantity) if r.quantity is not None else None,
                           "amount_paise": r.amount_paise, "source_row": r.source_row} for r in stale] if not has_errors else []
    if has_errors:
        tr["pruning"] = (f"skipped: the file has row errors, so {len(stale)} previously imported row(s) not in "
                         "this file were kept" if stale else "skipped (file has row errors)")
    else:
        for row in stale:
            changelog.record_txn(session, "TXN_REMOVED", "import", row, row.instrument.symbol or row.instrument.name)
            session.delete(row)
        tr["removed"] = len(stale)
        tr["pruning"] = f"removed {len(stale)} previously imported row(s) no longer in the file"
    session.flush()
    for row in added:
        changelog.record_txn(session, "TXN_ADDED", "import", row, row.instrument.symbol or row.instrument.name)

    # fallback: only for instruments with no transaction rows
    from sqlalchemy import delete as sa_delete

    from app.models import SectorAssignment

    # Excel-origin sectors mirror the current file (manual and NSE assignments are untouched)
    session.execute(sa_delete(SectorAssignment).where(SectorAssignment.origin == "EXCEL"))
    with_txns = repo.instrument_ids_with_transactions(session)
    fb_rows = []
    for f in fallbacks:
        inst = _instrument(session, cache, report, f.instrument_key, f.asset_type, f.name, f.symbol, f.platform)
        used = inst.id not in with_txns
        if f.sector:
            from app.services.classification import upsert_assignment

            upsert_assignment(session, inst.id, "EXCEL", f.sector, "Excel holdings_fallback.sector", f.as_of_date,
                              f"row {f.row}")
        if used:
            report["holdings_fallback"]["used"].append(f.name)
        fb_rows.append(FallbackHolding(
            instrument_id=inst.id, as_of_date=f.as_of_date, platform=f.platform, quantity=f.quantity,
            invested_paise=f.invested_paise, monthly_plan_paise=f.monthly_plan_paise, sector=f.sector,
            notes=f.notes, used=used, source_row=f.row,
        ))
    report["holdings_fallback"]["ignored"] = len(fb_rows) - len(report["holdings_fallback"]["used"])
    repo.replace_fallback(session, fb_rows)

    # monthly plan: replaced wholesale; MF rows linked to instruments by normalised name
    mf_by_name = {k.removeprefix("MF:"): i for k, i in cache.items() if k.startswith("MF:")}
    plan_rows = []
    for p in plans:
        inst = mf_by_name.get(normalise_name(p.name)) if not p.is_bucket else None
        if not p.is_bucket and inst is None:
            report["monthly_plan"]["warnings"].append(f"row {p.row}: '{p.name}' does not match any held fund")
        plan_rows.append(MonthlyPlanItem(
            platform=p.platform, asset_type=p.asset_type, name=p.name, monthly_amount_paise=p.monthly_amount_paise,
            allocation_fraction=p.allocation_fraction, instrument_id=inst.id if inst else None,
            is_bucket=p.is_bucket, source_row=p.row,
        ))
    repo.replace_plan(session, plan_rows)
    report["instruments"]["existing"] = existing_count
    report["changes_logged"] = changelog.record_position_changes(session, positions_before, "import")


def _finish(session: Session, report: dict) -> dict:
    tr = report["transactions"]
    report["ok"] = not report["fatal"] and not tr["errors"]
    session.add(ImportRun(
        file_name=report["file"], file_sha256=report["sha256"], ok=report["ok"],
        rows_loaded=tr["loaded"], rows_skipped=len(tr["skipped"]), rows_errors=len(tr["errors"]), report=report,
    ))
    session.flush()
    return report


def format_report(report: dict) -> str:
    tr, fb, mp = report["transactions"], report["holdings_fallback"], report["monthly_plan"]
    lines = [f"Import report — {report['file']}", "=" * 60]
    if report["fatal"]:
        lines += ["FATAL — nothing was imported:"] + [f"  - {m}" for m in report["fatal"]]
        return "\n".join(lines)
    by = ", ".join(f"{k} {v}" for k, v in tr["by_asset_type"].items())
    lines += [
        "transactions:",
        f"  rows in sheet : {tr['rows_in_sheet']}",
        f"  loaded        : {tr['loaded']}  ({by})",
        f"    new         : {tr['inserted']}",
        f"    unchanged   : {tr['unchanged']}" + (f" (notes updated on {tr['updated_notes']})" if tr["updated_notes"] else ""),
        f"  in-app dupes  : {len(tr['matched_ui'])} (already entered in the app; not duplicated)",
        f"  skipped       : {len(tr['skipped'])}",
        f"  errors        : {len(tr['errors'])}",
        f"  pruning       : {tr['pruning']}",
    ]
    lines += [f"    skipped row {s['row']}: {s['reason']}" for s in tr["skipped"]]
    lines += [f"    ERROR row {e['row']} [{e['field'] or '-'}]: {e['message']}" for e in tr["errors"]]
    lines += [f"    warning row {w['row']}: {w['message']}" for w in tr["warnings"]]
    lines += [
        "holdings_fallback:",
        f"  rows: {fb['rows']}, used: {len(fb['used'])}, ignored (instrument has transactions): {fb['ignored']}",
    ]
    lines += [f"    used for: {n}" for n in fb["used"]]
    lines += [f"    ERROR row {e['row']}: {e['message']}" for e in fb["errors"]]
    lines += [
        "monthly_plan:",
        f"  items: {mp['rows']}, monthly total: {format_inr(mp['total_paise'])}, allocation sum: {mp['fraction_sum']}",
    ]
    lines += [f"    warning: {w}" for w in mp["warnings"]]
    lines += [f"    ERROR row {e['row']}: {e['message']}" for e in mp["errors"]]
    created = report["instruments"]["created"]
    lines += [f"instruments: {len(created)} created, {report['instruments']['existing']} already known"]
    lines += ["", "RESULT: " + ("OK" if report["ok"] else "COMPLETED WITH ERRORS")]
    return "\n".join(lines)



def preview_workbook(session: Session, path: str | Path, file_name: str | None = None) -> dict:
    """Dry run: the real importer inside a SAVEPOINT that is always rolled back. Nothing is saved."""
    sp = session.begin_nested()
    try:
        report = import_workbook(session, path, file_name)
    finally:
        sp.rollback()
    session.expire_all()
    report["preview"] = True
    return report
