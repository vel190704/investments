import io
from decimal import Decimal as D

import openpyxl
import pytest
from sqlalchemy import func, select

from tests.conftest import make_user
from app.calc.positions import TxnInput, replay
from app.importer.run import import_workbook
from app.models import ChangeLogEntry, Instrument, Transaction
from app.services.export import build_workbook
from app.services.transactions import TxnValidationError, add_ui_transaction, delete_ui_transaction
from tests.conftest import TXN_HEADER, base_txn_rows, write_workbook


def count(s, model, *where):
    q = select(func.count()).select_from(model)
    for w in where:
        q = q.where(w)
    return s.scalar(q)


def inst(s, symbol=None, name_like=None):
    q = select(Instrument)
    q = q.where(Instrument.symbol == symbol) if symbol else q.where(Instrument.name.like(name_like))
    return s.scalar(q)


def kotak_sip(s):
    k = inst(s, name_like="Kotak%")
    return add_ui_transaction(s, {"instrument_id": k.id, "date": "2026-11-05", "type": "SIP",
                                  "quantity": "49.396", "price": "20.2442", "amount": "1000",
                                  "platform": "Zerodha Coin"})


def test_add_mf_sip_amount_authoritative(session, workbook):
    import_workbook(session, workbook)
    r = kotak_sip(session)
    t = session.get(Transaction, r["id"])
    assert t.source == "ui" and t.row_key.startswith("ui:") and t.amount_paise == 100000
    assert t.quantity == D("49.396") and r["warnings"] == []  # 49.396 × 20.2442 ≈ 999.98, within tolerance
    r2 = add_ui_transaction(session, {"instrument_id": t.instrument_id, "date": "2026-12-05", "type": "SIP",
                                      "quantity": "40", "price": "20", "amount": "1000"})
    assert any("amount is used" in w for w in r2["warnings"])


def test_ui_validation_matches_excel_rules(session, workbook):
    import_workbook(session, workbook)
    lt = inst(session, "LT")
    bad = [
        ({"type": "OPENING"}, "type"),
        ({"type": "HOLD"}, "type"),
        ({"date": ""}, "date"),
        ({"date": "05/11/2026"}, "date"),
        ({"quantity": "-1"}, "quantity"),
        ({"amount": "abc"}, "amount"),
        ({"type": "SELL", "quantity": "2"}, "quantity"),  # only 1 share held
    ]
    for override, field in bad:
        data = {"instrument_id": lt.id, "date": "2026-11-05", "type": "BUY", "quantity": "1", "price": "3700",
                "amount": "3700", **override}
        with pytest.raises(TxnValidationError) as e:
            add_ui_transaction(session, data)
        assert any(err["field"] == field for err in e.value.errors), override


def test_new_instrument_from_ui(session, workbook):
    import_workbook(session, workbook)
    add_ui_transaction(session, {"asset_type": "STOCK", "name": "Infosys", "symbol": "infy", "date": "2026-11-05",
                                 "type": "BUY", "quantity": "2", "price": "1500", "amount": "3000",
                                 "platform": "Groww"})
    assert inst(session, "INFY").instrument_key == "EQ:INFY"
    with pytest.raises(TxnValidationError):  # stocks need a symbol
        add_ui_transaction(session, {"asset_type": "STOCK", "name": "X", "date": "2026-11-05", "type": "BUY",
                                     "quantity": "1", "amount": "10"})


def test_reimport_keeps_ui_transaction(session, workbook):
    import_workbook(session, workbook)
    r = kotak_sip(session)
    rep = import_workbook(session, workbook)
    assert rep["ok"] and rep["transactions"]["removed"] == 0
    assert session.get(Transaction, r["id"]) is not None
    # even when Excel rows are pruned, ui rows survive
    rows = base_txn_rows()[:-1]
    import_workbook(session, write_workbook(workbook.parent / "fewer.xlsx", txn_rows=rows))
    assert session.get(Transaction, r["id"]) is not None
    assert count(session, Transaction, Transaction.source == "excel") == 18


def test_reimport_never_deletes_change_log(session, workbook):
    import_workbook(session, workbook)
    kotak_sip(session)
    n = count(session, ChangeLogEntry)
    import_workbook(session, write_workbook(workbook.parent / "fewer.xlsx", txn_rows=base_txn_rows()[:-1]))
    assert count(session, ChangeLogEntry) > n


def test_excel_duplicate_of_ui_row_is_not_inserted(session, workbook, tmp_path):
    import_workbook(session, workbook)
    kotak_sip(session)
    row = ["2026-11-05", "SIP", "MF", "Kotak Nifty Next 50 Index Fund - Direct - Growth", None, None, None,
           49.396, 20.2442, 1000, 0, "Zerodha Coin", "SIP typed into Excel too"]
    rep = import_workbook(session, write_workbook(tmp_path / "dup.xlsx", txn_rows=base_txn_rows() + [row]))
    assert rep["transactions"]["inserted"] == 0 and len(rep["transactions"]["matched_ui"]) == 1
    assert count(session, Transaction) == 20


def test_export_roundtrip(session, workbook, tmp_path):
    import_workbook(session, workbook)
    kotak_sip(session)
    data = build_workbook(session)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["README", "transactions", "holdings_fallback", "monthly_plan"]
    assert [c.value for c in wb["transactions"][1]] == TXN_HEADER
    assert wb["transactions"].max_row == 21
    path = tmp_path / "export.xlsx"
    path.write_bytes(data)
    # same DB: nothing new, ui row recognised, nothing removed
    rep = import_workbook(session, path)
    assert rep["ok"] and rep["transactions"]["inserted"] == 0 and rep["transactions"]["removed"] == 0
    assert count(session, Transaction) == 20
    assert rep["monthly_plan"]["warnings"] == []


def test_export_imports_into_empty_db_with_identical_positions(session, workbook, tmp_path):
    import_workbook(session, workbook)
    kotak_sip(session)

    def positions(s):
        out = {}
        for i in s.scalars(select(Instrument)):
            p = replay([TxnInput(t.date, t.type, t.quantity, t.amount_paise, t.charges_paise)
                        for t in s.scalars(select(Transaction).where(Transaction.instrument_id == i.id))])
            out[i.instrument_key] = (p.quantity, p.invested_paise)
        return out

    expected = positions(session)
    path = tmp_path / "export.xlsx"
    path.write_bytes(build_workbook(session))
    from app import db

    db.configure(f"sqlite:///{tmp_path / 'fresh.db'}")
    db.init_db()
    with db.session_scope(user_id=make_user()) as s2:
        rep = import_workbook(s2, path)
        assert rep["ok"] and rep["transactions"]["loaded"] == 20
        assert positions(s2) == expected


def test_delete_only_ui_rows_and_protect_later_sells(session, workbook):
    import_workbook(session, workbook)
    lt = inst(session, "LT")
    buy = add_ui_transaction(session, {"instrument_id": lt.id, "date": "2026-11-05", "type": "BUY",
                                       "quantity": "1", "price": "3700", "amount": "3700"})
    add_ui_transaction(session, {"instrument_id": lt.id, "date": "2026-11-20", "type": "SELL",
                                 "quantity": "2", "price": "3800", "amount": "7600"})
    with pytest.raises(TxnValidationError, match="later sale"):
        delete_ui_transaction(session, buy["id"])
    excel_row = session.scalar(select(Transaction).where(Transaction.source == "excel"))
    with pytest.raises(TxnValidationError, match="workbook"):
        delete_ui_transaction(session, excel_row.id)


def test_change_log_records_ui_add_and_position_change(session, workbook):
    import_workbook(session, workbook)
    lt = inst(session, "LT")
    add_ui_transaction(session, {"instrument_id": lt.id, "date": "2026-11-05", "type": "BUY",
                                 "quantity": "1", "price": "3600", "amount": "3600"})
    kinds = [e.kind for e in session.scalars(select(ChangeLogEntry).where(ChangeLogEntry.cause == "ui"))]
    assert kinds == ["TXN_ADDED", "QTY_CHANGED", "AVG_COST_CHANGED"]
    q = session.scalar(select(ChangeLogEntry).where(ChangeLogEntry.kind == "AVG_COST_CHANGED"))
    assert q.summary == "LT: average cost ₹3768.00 → ₹3684.00"


def test_import_oversell_check_includes_ui_rows(session, workbook, tmp_path):
    import_workbook(session, workbook)
    lt = inst(session, "LT")
    add_ui_transaction(session, {"instrument_id": lt.id, "date": "2026-11-05", "type": "BUY",
                                 "quantity": "1", "price": "3600", "amount": "3600"})
    sell = ["2026-11-20", "SELL", "STOCK", "Larsen & Toubro", "LT", None, None, 2, 3800, 7600, 0, "Groww", None]
    rep = import_workbook(session, write_workbook(tmp_path / "s.xlsx", txn_rows=base_txn_rows() + [sell]))
    assert rep["ok"] and rep["transactions"]["inserted"] == 1  # 1 (excel) + 1 (ui) held
