import os
from datetime import date
from pathlib import Path

import openpyxl
import pytest

os.environ["SCHEDULER_ENABLED"] = "false"
import tempfile  # noqa: E402
os.environ.setdefault("LOG_DIR", tempfile.mkdtemp(prefix="investments-test-logs-"))  # never the repo logs/

from app import db  # noqa: E402

TXN_HEADER = ["date", "type", "asset_type", "name", "symbol", "isin", "scheme_code", "quantity", "price",
              "amount", "charges", "platform", "notes"]

# Mirrors the user's workbook (19 OPENING rows: 14 stocks, 2 ETFs, 3 funds).
STOCKS = [
    ("Larsen & Toubro", "LT", "STOCK", 3768), ("Reliance Industries", "RELIANCE", "STOCK", 1193.8),
    ("State Bank of India", "SBIN", "STOCK", 963.95), ("HDFC Bank", "HDFCBANK", "STOCK", 718.35),
    ("KPIT Technologies", "KPITTECH", "STOCK", 529.95), ("Nippon India ETF Hang Seng BeES", "HNGSNGBEES", "ETF", 459),
    ("Coal India", "COALINDIA", "STOCK", 428.25), ("Bharat Electronics", "BEL", "STOCK", 398.95),
    ("Biocon", "BIOCON", "STOCK", 392.45), ("Indus Towers", "INDUSTOWER", "STOCK", 383.8),
    ("Tata Power", "TATAPOWER", "STOCK", 368.8), ("ITC", "ITC", "STOCK", 266.4),
    ("Mirae Asset NYSE FANG+ ETF", "MAFANG", "ETF", 235.7), ("MosChip Technologies", "MOSCHIP", "STOCK", 204.07),
    ("Tata Steel", "TATASTEEL", "STOCK", 183.55), ("Ashok Leyland", "ASHOKLEY", "STOCK", 162),
]
FUNDS = [
    ("HDFC Mid Cap Fund - Direct - Growth", 10.945, 228.4148, 2500),
    ("ICICI Prudential Nifty 50 Index Fund - Direct - Growth", 16.359, 244.5137, 4000),
    ("Kotak Nifty Next 50 Index Fund - Direct - Growth", 72.377, 20.7248, 1500),
]
PLAN = [
    ("Groww", "STOCK/ETF", "Direct stocks & ETFs", 6000, 0.6),
    ("Zerodha Coin", "MF", "HDFC Mid Cap Fund - Direct - Growth", 1500, 0.15),
    ("Zerodha Coin", "MF", "ICICI Prudential Nifty 50 Index Fund - Direct - Growth", 1500, 0.15),
    ("Zerodha Coin", "MF", "Kotak Nifty Next 50 Index Fund - Direct - Growth", 1000, 0.1),
    ("TOTAL", "ALL", "Total monthly investment", 10000, 1),
]


def base_txn_rows():
    rows = [["2026-10-03", "OPENING", at, n, sym, None, None, 1, p, p, 0, "Groww", "Opening balance"]
            for n, sym, at, p in STOCKS]
    rows += [["2026-10-03", "OPENING", "MF", n, None, None, None, q, p, a, 0, "Zerodha Coin", "Opening balance"]
             for n, q, p, a in FUNDS]
    return rows


def write_workbook(path: Path, txn_rows=None, fallback_rows=None, plan_rows=None, sheets=None) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "README"
    ws.append(["ignored"])
    sheets = sheets or ["transactions", "holdings_fallback", "monthly_plan"]
    if "transactions" in sheets:
        t = wb.create_sheet("transactions")
        t.append(TXN_HEADER)
        for r in (base_txn_rows() if txn_rows is None else txn_rows):
            t.append(r)
    if "holdings_fallback" in sheets:
        f = wb.create_sheet("holdings_fallback")
        f.append(["as_of_date", "platform", "asset_type", "name", "symbol", "quantity", "invested_amount",
                  "monthly_plan", "sector", "notes"])
        for r in (fallback_rows or []):
            f.append(r)
    if "monthly_plan" in sheets:
        p = wb.create_sheet("monthly_plan")
        p.append(["platform", "asset_type", "name", "monthly_amount", "allocation_percent"])
        for r in (PLAN if plan_rows is None else plan_rows):
            p.append(list(r))
    wb.save(path)
    return path


@pytest.fixture(autouse=True)
def isolate_from_dotenv(monkeypatch):
    """Tests never see the developer's real secrets or model choice from .env."""
    from app.config import get_settings

    st = get_settings()
    monkeypatch.setattr(st, "anthropic_api_key", "")
    if hasattr(st, "claude_model"):
        monkeypatch.setattr(st, "claude_model", "")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)


def make_user(email: str = "owner@localhost") -> int:
    """Create (or find) a user the way dev auth does; the default is the dev user, i.e. the API's user in tests."""
    from app.auth import resolve_user

    return resolve_user(email, None, dev=True).id


@pytest.fixture
def session(tmp_path):
    """A session bound to the default (dev/admin) user: the same user TestClient requests act as."""
    db.configure(f"sqlite:///{tmp_path / 'test.db'}")
    db.init_db()
    uid = make_user()
    with db.session_scope(user_id=uid) as s:
        yield s


@pytest.fixture
def workbook(tmp_path):
    return write_workbook(tmp_path / "book.xlsx")


VALUATION_DATE = date(2026, 10, 3)


@pytest.fixture
def pg_url():
    """A fresh, empty Postgres database (needs TEST_POSTGRES_URL, an admin URL such as
    postgresql+psycopg://postgres:pw@127.0.0.1:55432/postgres). Skipped when not set."""
    import uuid

    from sqlalchemy import create_engine, text
    from sqlalchemy.engine import make_url

    admin = os.environ.get("TEST_POSTGRES_URL")
    if not admin:
        pytest.skip("TEST_POSTGRES_URL not set")
    name = f"t_{uuid.uuid4().hex[:10]}"
    eng = create_engine(admin, isolation_level="AUTOCOMMIT")
    with eng.connect() as c:
        c.execute(text(f'CREATE DATABASE "{name}"'))
    url = make_url(admin).set(database=name).render_as_string(hide_password=False)
    try:
        yield url
    finally:
        db.configure("sqlite://")  # release the pool before dropping
        with eng.connect() as c:
            c.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
        eng.dispose()


_PG_TEMPLATE: dict = {}


@pytest.fixture(autouse=True)
def all_on_postgres(monkeypatch):
    """TEST_ALL_ON_POSTGRES=1 (+ TEST_POSTGRES_URL): every db.configure('sqlite:///…') in a test gets a fresh
    Postgres database instead, cloned from a template migrated once. Finds Postgres-only bugs SQLite hides."""
    if os.environ.get("TEST_ALL_ON_POSTGRES") != "1" or not os.environ.get("TEST_POSTGRES_URL"):
        yield
        return
    import uuid

    from sqlalchemy import create_engine, text
    from sqlalchemy.engine import make_url

    admin_url = os.environ["TEST_POSTGRES_URL"]
    admin = create_engine(admin_url, isolation_level="AUTOCOMMIT")
    if "name" not in _PG_TEMPLATE:
        name = f"tmpl_{uuid.uuid4().hex[:8]}"
        with admin.connect() as c:
            c.execute(text(f'CREATE DATABASE "{name}"'))
        url = make_url(admin_url).set(database=name).render_as_string(hide_password=False)
        db.configure(url)
        db.init_db()
        db.get_engine().dispose()
        _PG_TEMPLATE["name"] = name
    created: list[str] = []
    real_configure = db.configure

    def configure(url=None):
        if url and url.startswith("sqlite:///"):
            name = f"t_{uuid.uuid4().hex[:10]}"
            with admin.connect() as c:
                c.execute(text(f'CREATE DATABASE "{name}" TEMPLATE "{_PG_TEMPLATE["name"]}"'))
            created.append(name)
            url = make_url(admin_url).set(database=name).render_as_string(hide_password=False)
        return real_configure(url)

    monkeypatch.setattr(db, "configure", configure)
    yield
    real_configure("sqlite://")
    with admin.connect() as c:
        for name in created:
            c.execute(text(f'DROP DATABASE IF EXISTS "{name}" WITH (FORCE)'))
    admin.dispose()
