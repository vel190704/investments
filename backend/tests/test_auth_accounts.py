"""Auth (Supabase JWT, invite-only, dev mode guard) and account management (admin, self-service, own key)."""

import time

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import ec
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from app import auth, db
from app.config import get_settings
from tests.conftest import make_user

SUPA = "https://proj.supabase.co"
SECRET = "test-hs256-secret-0123456789abcdef0123"


def token(email="alice@example.com", sub="11111111-1111-1111-1111-111111111111", exp_in=3600, aud="authenticated",
          iss=f"{SUPA}/auth/v1", key=SECRET, alg="HS256", kid=None):
    claims = {"sub": sub, "email": email, "aud": aud, "iss": iss, "exp": int(time.time()) + exp_in, "role": "authenticated"}
    return jwt.encode(claims, key, algorithm=alg, headers={"kid": kid} if kid else None)


@pytest.fixture
def supa(tmp_path, monkeypatch):
    st = get_settings()
    for k, v in {"dev_auth": "", "supabase_url": SUPA, "supabase_jwt_secret": SECRET,
                 "admin_emails": "owner@example.com", "env": "production"}.items():
        monkeypatch.setattr(st, k, v)
    db.configure(f"sqlite:///{tmp_path / 'a.db'}")
    # these tests are about tokens; the database-role guard has its own tests (test_postgres.py). Under
    # TEST_ALL_ON_POSTGRES the test role is a superuser, which that guard (correctly) refuses in production.
    monkeypatch.setattr(db, "check_rls_role", lambda: None)
    # likewise the "production needs PostgreSQL" guard (test_database_config.py): these tests run on SQLite
    monkeypatch.setattr(db, "check_database_config", lambda url=None: {"dialect": "sqlite", "driver": "-", "host": "-"})
    with TestClient(app_()) as c:
        yield c


def app_():
    from app.main import app

    return app


def H(t):
    return {"Authorization": f"Bearer {t}"}


def test_requests_need_a_valid_token(supa):
    assert supa.get("/api/me").status_code == 401
    assert supa.get("/api/me", headers=H("garbage")).status_code == 401
    assert supa.get("/api/me", headers=H(token(exp_in=-10))).json()["detail"] == "session expired: sign in again"
    assert supa.get("/api/me", headers=H(token(aud="anon"))).status_code == 401
    assert supa.get("/api/me", headers=H(token(iss="https://evil.example/auth/v1"))).status_code == 401
    assert supa.get("/api/me", headers=H(token(key="another-secret-0123456789abcdef0000"))).status_code == 401
    assert supa.get("/api/health").status_code == 200  # public


def test_invite_only_and_admin_bootstrap(supa):
    r = supa.get("/api/me", headers=H(token("stranger@example.com")))
    assert r.status_code == 403 and "not invited" in r.json()["detail"]
    owner = supa.get("/api/me", headers=H(token("owner@example.com", sub="o-1")))
    assert owner.status_code == 200 and owner.json()["is_admin"] is True  # ADMIN_EMAILS
    oh = H(token("owner@example.com", sub="o-1"))
    assert supa.post("/api/admin/invites", headers=oh, json={"email": "Friend@Example.com", "send_email": False}).status_code == 201
    friend = supa.get("/api/me", headers=H(token("friend@example.com", sub="f-1")))
    assert friend.status_code == 200 and friend.json()["role"] == "user" and friend.json()["ai"]["ai_enabled_by_admin"] is False
    assert supa.get("/api/admin/users", headers=H(token("friend@example.com", sub="f-1"))).status_code == 403
    assert supa.delete("/api/admin/invites/friend@example.com", headers=oh).status_code == 200
    r = supa.get("/api/me", headers=H(token("friend@example.com", sub="f-1")))
    assert r.status_code == 403 and "invitation was removed" in r.json()["detail"]


@pytest.mark.parametrize("env, dev_auth, supabase, outcome", [
    ("development", "1", "", "dev"),          # the only ways to get the fake local user
    ("test", "1", "", "dev"),
    ("production", "1", SUPA, "refuse"),      # DEV_AUTH outside development/test refuses to start
    ("preview", "1", SUPA, "refuse"),         # preview deployments are real deployments
    ("", "1", SUPA, "refuse"),                # unset ENV = production
    ("staging", "1", SUPA, "refuse"),         # unknown ENV = production
    ("development", "", SUPA, "auth"),        # DEV_AUTH not set: sign-in required even locally
    ("development", "true", SUPA, "auth"),    # only the exact value "1" enables it
    ("", "", SUPA, "auth"),
    ("", "", "", "refuse"),                   # no auth configured at all: refuse to start
    ("development", "", "", "refuse"),
])
def test_authentication_fails_closed(monkeypatch, env, dev_auth, supabase, outcome):
    st = get_settings()
    monkeypatch.setattr(st, "env", env)
    monkeypatch.setattr(st, "dev_auth", dev_auth)
    monkeypatch.setattr(st, "supabase_url", supabase)
    if outcome == "refuse":
        with pytest.raises(RuntimeError):
            auth.check_startup_config()
        return
    auth.check_startup_config()
    assert auth.dev_auth_active() is (outcome == "dev")
    assert auth.effective_env() == (env if env in auth.KNOWN_ENVS else "production")


def test_without_dev_auth_requests_need_a_token(tmp_path, monkeypatch):
    st = get_settings()
    monkeypatch.setattr(st, "env", "development")
    monkeypatch.setattr(st, "dev_auth", "")
    monkeypatch.setattr(st, "supabase_url", SUPA)
    monkeypatch.setattr(st, "supabase_jwt_secret", SECRET)
    db.configure(f"sqlite:///{tmp_path / 'fc.db'}")
    with TestClient(app_()) as c:
        assert c.get("/api/me").status_code == 401
        assert c.get("/api/me", headers={"X-Dev-User": "owner@localhost"}).status_code == 401  # header ignored
        assert c.get("/api/health").status_code == 200


def test_asymmetric_tokens_via_jwks(monkeypatch):
    key = ec.generate_private_key(ec.SECP256R1())
    st = get_settings()
    monkeypatch.setattr(st, "supabase_url", SUPA)
    monkeypatch.setattr(st, "supabase_jwt_secret", "")

    class FakeJWKS:
        def get_signing_key_from_jwt(self, t):
            assert jwt.get_unverified_header(t)["kid"] == "k1"
            return type("K", (), {"key": key.public_key()})()

    monkeypatch.setattr(auth, "_jwks_client", lambda url: FakeJWKS() if url == f"{SUPA}/auth/v1/.well-known/jwks.json" else None)
    claims = auth.verify_token(token(key=key, alg="ES256", kid="k1"))
    assert claims["email"] == "alice@example.com"
    with pytest.raises(auth.AuthError):  # HS256 is refused when no legacy secret is configured
        auth.verify_token(token())
    other = ec.generate_private_key(ec.SECP256R1())
    with pytest.raises(auth.AuthError):
        auth.verify_token(token(key=other, alg="ES256", kid="k1"))


def test_admin_can_delete_a_users_data_completely(tmp_path, workbook):
    from app.importer.run import import_workbook
    from app.models import AllowedEmail, Transaction, User
    from app.services.accounts import delete_user_data, list_users

    db.configure(f"sqlite:///{tmp_path / 'd.db'}")
    db.init_db()
    a, b = make_user("alice@example.com"), make_user("bob@example.com")
    for uid in (a, b):
        with db.session_scope(user_id=uid) as s:
            assert import_workbook(s, workbook)["ok"]
    with db.session_scope(system=True) as s:
        s.add(AllowedEmail(email="bob@example.com"))
    with db.session_scope(system=True) as s:
        users = {u["email"]: u for u in list_users(s)}
        assert users["bob@example.com"]["rows"]["transactions"] == 19
        res = delete_user_data(s, b)
        assert res["account_removed"] and res["deleted_rows"]["transactions"] == 19
    with db.session_scope(system=True) as s:
        assert s.get(User, b) is None and s.get(AllowedEmail, "bob@example.com") is None
        assert s.scalar(select(func.count()).select_from(Transaction)) == 19  # Alice's are untouched
        from app.services.accounts import _user_tables
        for t in _user_tables():
            assert s.scalar(select(func.count()).select_from(t).where(t.c.user_id == b)) == 0, t.name


def test_delete_my_account_needs_my_email(tmp_path, workbook):
    db.configure(f"sqlite:///{tmp_path / 'me.db'}")
    with TestClient(app_()) as c:
        A = {"X-Dev-User": "alice@example.com"}
        c.post("/api/import", headers=A, files={"file": ("w.xlsx", workbook.read_bytes())})
        assert c.request("DELETE", "/api/me", headers=A, json={"confirm": "bob@example.com"}).status_code == 422
        r = c.request("DELETE", "/api/me", headers=A, json={"confirm": "Alice@Example.com"})
        assert r.status_code == 200 and r.json()["deleted_rows"]["transactions"] == 19
        assert c.get("/api/me", headers=A).json()["has_data"] is False  # dev mode re-creates an empty account


def test_own_key_is_encrypted_and_never_returned(tmp_path, monkeypatch):
    from cryptography.fernet import Fernet

    from app.models import User

    monkeypatch.setattr(get_settings(), "field_encryption_key", Fernet.generate_key().decode())
    db.configure(f"sqlite:///{tmp_path / 'k.db'}")
    key = "sk-ant-api03-" + "Z" * 40
    with TestClient(app_()) as c:
        A = {"X-Dev-User": "alice@example.com"}
        r = c.put("/api/me/ai-key", headers=A, json={"key": key})
        assert r.status_code == 200 and r.json() == {"own_key": True}
        assert key not in c.get("/api/me", headers=A).text
        assert c.put("/api/me/ai-key", headers=A, json={"key": "nope"}).status_code == 422
    with db.session_scope(system=True) as s:
        stored = s.scalar(select(User.own_key_encrypted).where(User.email == "alice@example.com"))
        assert stored and key not in stored
        from app.crypto import decrypt
        assert decrypt(stored) == key


def test_non_admin_ai_is_off_until_enabled(session, monkeypatch):
    from app.models import User
    from app.services import analysis as an
    from app.services import settings as st

    monkeypatch.setattr(get_settings(), "anthropic_api_key", "sk-test")
    friend = make_user("friend@example.com")
    with db.session_scope(user_id=friend) as s:
        st.put(s, "llm", {**st.get(s, "llm"), "outbound_enabled": True})
        assert "not enabled for your account" in an.gate(s)
    with db.session_scope(system=True) as s:
        s.get(User, friend).ai_enabled = True
    with db.session_scope(user_id=friend) as s:
        assert an.gate(s) is None
        assert an.usage(s)["cap"] == 10  # admin's default per-user cap


def test_global_ai_cap_counts_platform_key_calls_of_everyone(session, monkeypatch):
    from app.models import LlmCall
    from app.services import analysis as an
    from app.services import settings as st

    monkeypatch.setattr(get_settings(), "anthropic_api_key", "sk-test")
    st.put(session, "llm", {**st.get(session, "llm"), "outbound_enabled": True})
    st.put(session, "llm_platform", {"global_monthly_cap": 2})
    session.commit()
    other = make_user("other@example.com")
    with db.session_scope(user_id=other) as s:
        s.add_all([LlmCall(purpose="stock", model="m", ok=True), LlmCall(purpose="stock", model="m", ok=True)])
    assert "monthly AI budget is used up (2/2)" in an.gate(session)


def test_template_has_headers_and_no_data_rows(tmp_path):
    import io

    import openpyxl

    from app.importer.run import import_workbook
    from app.services.template import build_template

    wb = openpyxl.load_workbook(io.BytesIO(build_template()))
    assert wb.sheetnames == ["README", "transactions", "holdings_fallback", "monthly_plan"]
    assert [c.value for c in wb["transactions"][1]][:4] == ["date", "type", "asset_type", "name"]
    assert all(wb[n].max_row == 1 for n in ("transactions", "holdings_fallback", "monthly_plan"))
    p = tmp_path / "t.xlsx"
    p.write_bytes(build_template())
    db.configure(f"sqlite:///{tmp_path / 't.db'}")
    db.init_db()
    with db.session_scope(user_id=make_user()) as s:
        rep = import_workbook(s, p)
    assert rep["transactions"]["loaded"] == 0  # the blank template imports cleanly with nothing in it


def test_changed_encryption_key_blocks_own_key_use_without_falling_back(session, monkeypatch):
    """If FIELD_ENCRYPTION_KEY changes, stored own keys cannot be read: AI is blocked with a clear message (the
    person re-enters their key); it never silently switches to the app's shared key."""
    from cryptography.fernet import Fernet

    from app.crypto import encrypt
    from app.models import User
    from app.services import analysis as an
    from app.services import settings as st

    st_ = get_settings()
    monkeypatch.setattr(st_, "anthropic_api_key", "sk-platform-key-not-to-be-used")
    monkeypatch.setattr(st_, "field_encryption_key", Fernet.generate_key().decode())
    st.put(session, "llm", {**st.get(session, "llm"), "outbound_enabled": True})
    session.get(User, session.info["user_id"]).own_key_encrypted = encrypt("sk-ant-api03-" + "Y" * 40)
    session.flush()
    assert an.gate(session) is None and an._api_key(session) == ("sk-ant-api03-" + "Y" * 40, True)
    monkeypatch.setattr(st_, "field_encryption_key", Fernet.generate_key().decode())  # the secret changed
    assert "cannot be decrypted" in an.gate(session)


def test_simple_mode_defaults_by_role_and_can_be_chosen(tmp_path):
    db.configure(f"sqlite:///{tmp_path / 'ui.db'}")
    with TestClient(app_()) as c:
        member, admin = {"X-Dev-User": "friend@example.com"}, {}  # the default dev user is the admin
        assert c.get("/api/me", headers=member).json()["simple_mode"] is True
        assert c.get("/api/me", headers=admin).json()["simple_mode"] is False
        assert c.put("/api/settings/ui", headers=member, json={"simple_mode": False}).status_code == 200
        me = c.get("/api/me", headers=member).json()
        assert me["simple_mode"] is False and me["simple_mode_chosen"] is False
        assert c.get("/api/me", headers={"X-Dev-User": "other@example.com"}).json()["simple_mode"] is True  # per user
        assert c.put("/api/settings/ui", headers=member, json={"simple_mode": "yes"}).status_code == 422
        assert c.put("/api/settings/ui", headers=member, json={"simple_mode": None}).status_code == 200
        assert c.get("/api/me", headers=member).json()["simple_mode"] is True  # back to the role default
